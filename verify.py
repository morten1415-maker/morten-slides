# -*- coding: utf-8 -*-
"""
Regressionstjek af slide-systemet
=================================
Der er ingen PowerPoint- eller browser-rendering i denne pipeline, så de faste
mål kan skride uden at nogen ser det. Dette script er værnet: det bygger
galleriet og kontrollerer de invarianter der før blev rettet i hånden deck
efter deck.

    python verify.py                # byg example.pptx og tjek alt
    python verify.py mit-deck.pptx  # tjek en bestemt fil

Exit-kode 0 = alt bestået, 1 = mindst én fejl.

Seks grupper:
  0. Referencemål       — at konstanterne stadig er dem der blev udmålt i de
                          to referencedeck (fanges ellers ikke, da gruppe 1
                          udleder sine forventninger fra netop konstanterne)
  1. Geometri i .pptx   — overskrift, grøn accent, logo, kilde-linje, forside
  2. Font-indlejring    — at IBM Plex faktisk ligger i filen og er hægtet rigtigt op
  3. HTML/PPTX-paritet  — at slides.css og build_pptx.py taler om samme mål
  4. Galleri-dækning    — at hver slide-type i Deck også vises i deck.example.py
  5. Oprydning          — ingen efterladte kicker-labels over titler
"""
from __future__ import annotations

import os
import re
import subprocess
import sys
import zipfile
import hashlib

HERE = os.path.dirname(os.path.abspath(__file__))
FONTS_DIR = os.path.join(HERE, "assets", "fonts")

# Canvas'et er præcis 96 px pr. tomme i begge akser:
#   1280 px / 13.333" = 96   og   720 px / 7.5" = 96
PX_PER_INCH = 96.0

_failures: list[str] = []


def check(label, cond, detail=""):
    if not cond:
        _failures.append(label)
    print("  %s %-54s %s" % ("OK  " if cond else "FEJL", label, detail))
    return bool(cond)


def near(a, b, tol):
    return abs(a - b) <= tol


def head(title):
    print("\n%s\n%s" % (title, "-" * len(title)))


# =================================================================== 0. REFERENCEMÅL
# Målene nedenfor er UDMÅLT i de to referencedeck Morten rettede i hånden
# ("Morten SKF.pptx" og "Hvem kan få standardrapport.pptx") og er systemets
# sandhed. De står som tal her — ikke som import fra build_pptx — netop så et
# uheldigt tryk på en konstant i build_pptx.py bliver fanget i stedet for
# stiltiende at flytte forventningen med sig. Ændrer du bevidst et mål, skal
# det ændres BÅDE her og i build_pptx.py.
REFERENCE_CM = {
    "MARGIN": 1.68,          # sikker margin til alle kanter
    "TITLE_TOP": 1.27,       # overskriftens tekstboks, top
    "ACCENT_X": 1.93,        # grøn accent, venstrekant (flush med titel-teksten)
    "ACCENT_Y": 3.26,        # grøn accent, top — fri af overskriftens underlængder
    "ACCENT_W": 1.78,        # grøn accent, bredde
    "ACCENT_H": 0.21,        # grøn accent, højde (6pt)
    "CONTENT_TOP": 4.32,     # indholdet starter
    "CONTENT_BOT": 16.76,    # indholdet slutter
    "SOURCE_Y": 17.40,       # kilde-/note-linje
    "LOGO_W": 3.94,          # logo, bredde
    "LOGO_INSET": 0.71,      # logo, ens luft til højre og bund
    "COVER_TOP": 4.32,       # forsidens kasse, top
    "COVER_W": 27.94,        # forsidens kasse, bredde
    "COVER_H": 10.16,        # forsidens kasse, højde med manchet
    "COVER_H_SHORT": 9.14,   # forsidens kasse, højde uden manchet
    "BIGNUM_TOP": 4.32,      # stort tal, blokkens top
}


def check_reference_measurements():
    import build_pptx as bp
    for name, want in sorted(REFERENCE_CM.items()):
        emu = getattr(bp, name, None)
        if emu is None:
            check("%s findes i build_pptx.py" % name, False)
            continue
        got = emu / 914400.0 * 2.54
        check("%-14s = %.2f cm" % (name, want), near(got, want, 0.02),
              "" if near(got, want, 0.02) else "er %.2f cm" % got)


# ======================================================================= 1. GEOMETRI
def check_pptx_geometry(path):
    from pptx import Presentation
    from pptx.util import Emu
    import build_pptx as bp

    def cm(v):
        return round(Emu(v).cm, 2)

    inch_cm = 2.54
    W_CM, H_CM = 13.333 * inch_cm, 7.5 * inch_cm

    exp_title = (round(bp.MARGIN / 914400 * inch_cm, 2),
                 round(bp.TITLE_TOP / 914400 * inch_cm, 2))
    exp_accent = (round(bp.ACCENT_X / 914400 * inch_cm, 2),
                  round(bp.ACCENT_Y / 914400 * inch_cm, 2))
    acc_size = (round(bp.ACCENT_W / 914400 * inch_cm, 2),
                round(bp.ACCENT_H / 914400 * inch_cm, 2))
    exp_src = (round(bp.MARGIN / 914400 * inch_cm, 2),
               round(bp.SOURCE_Y / 914400 * inch_cm, 2))
    logo_inset = bp.LOGO_INSET / 914400 * inch_cm
    logo_w = round(bp.LOGO_W / 914400 * inch_cm, 2)

    prs = Presentation(path)
    titles = accents = logos = sources = covers = 0
    bad = []

    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            l, t, w, h = cm(sh.left), cm(sh.top), cm(sh.width), cm(sh.height)

            # overskrift: tekstboks med topkant på TITLE_TOP
            if sh.shape_type == 17 and t == exp_title[1]:
                titles += 1
                if (l, t) != exp_title:
                    bad.append("slide %d: overskrift på %s, ventet %s" % (i, (l, t), exp_title))

            # grøn accent: rektangel i accent-størrelse
            if sh.name.startswith("Rectangle") and (w, h) == acc_size:
                accents += 1
                if (l, t) != exp_accent:
                    bad.append("slide %d: accent på %s, ventet %s" % (i, (l, t), exp_accent))

            # logo: billede i logo-bredde -> ens luft til højre og bund
            if sh.name.startswith("Picture") and w == logo_w:
                logos += 1
                r, b = W_CM - (l + w), H_CM - (t + h)
                if not near(r, b, 0.05):
                    bad.append("slide %d: logo-luft h=%.2f b=%.2f (skal være ens)" % (i, r, b))
                if not near(r, logo_inset, 0.05):
                    bad.append("slide %d: logo-luft %.2f cm, ventet %.2f" % (i, r, logo_inset))

            # kilde-linje
            if sh.has_text_frame and sh.text_frame.text.upper().startswith("KILDE"):
                sources += 1
                if (l, t) != exp_src:
                    bad.append("slide %d: kilde-linje på %s, ventet %s" % (i, (l, t), exp_src))

            # forsidens kasse
            if (sh.name.startswith("Rectangle")
                    and w == round(bp.COVER_W / 914400 * inch_cm, 2)):
                covers += 1
                exp_h = [round(bp.COVER_H / 914400 * inch_cm, 2),
                         round(bp.COVER_H_SHORT / 914400 * inch_cm, 2)]
                if (l, t) != (exp_title[0], round(bp.COVER_TOP / 914400 * inch_cm, 2)):
                    bad.append("slide %d: forside-kasse på %s" % (i, (l, t)))
                if h not in exp_h:
                    bad.append("slide %d: forside-kasse højde %s, ventet en af %s" % (i, h, exp_h))

    check("overskrifter på %s cm" % exp_title[1], titles > 0, "%d fundet" % titles)
    check("grøn accent: én pr. overskrift", accents == titles,
          "%d accenter / %d overskrifter" % (accents, titles))
    check("grøn accent på %s cm, flush med titlen" % (exp_accent,), not any("accent" in m for m in bad))
    check("logo i hjørnet med ens luft", logos > 0 and not any("logo" in m for m in bad),
          "%d slides" % logos)
    check("kilde-linjer på %s cm" % exp_src[1], not any("kilde" in m for m in bad),
          "%d fundet" % sources)
    check("forsidens kasse har referencemålene", covers > 0 and not any("forside" in m for m in bad),
          "%d slides" % covers)
    check("ingen geometri-afvigelser i alt", not bad, "%d afvigelser" % len(bad))
    for m in bad:
        print("       - " + m)
    return len(prs.slides._sldIdLst)


# ================================================================ 2. FONT-INDLEJRING
def check_font_embedding(path):
    from lxml import etree

    NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
          "pr": "http://schemas.openxmlformats.org/package/2006/relationships"}
    # CT_Presentation-rækkefølgen ifølge ECMA-376 — embeddedFontLst SKAL ligge her
    ORDER = ["sldMasterIdLst", "notesMasterIdLst", "handoutMasterIdLst", "sldIdLst",
             "sldSz", "notesSz", "smartTags", "embeddedFontLst", "custShowLst",
             "photoAlbum", "custDataLst", "kinsoku", "defaultTextStyle",
             "modifyVerifier", "extLst"]

    z = zipfile.ZipFile(path)
    names = z.namelist()

    ct = etree.fromstring(z.read("[Content_Types].xml"))
    defaults = {d.get("Extension").lower(): d.get("ContentType")
                for d in ct.findall("ct:Default", NS)}
    check("[Content_Types].xml: Default for fntdata",
          defaults.get("fntdata") == "application/x-fontdata",
          defaults.get("fntdata") or "mangler")

    parts = sorted(n for n in names if n.startswith("ppt/fonts/"))
    src = {}
    if os.path.isdir(FONTS_DIR):
        for f in sorted(os.listdir(FONTS_DIR)):
            if f.lower().endswith((".ttf", ".otf")):
                blob = open(os.path.join(FONTS_DIR, f), "rb").read()
                src[hashlib.sha256(blob).hexdigest()] = f
    check("fontfiler i assets/fonts/", len(src) >= 4, "%d filer" % len(src))
    check("OFL-licensen ligger hos fonten",
          os.path.exists(os.path.join(FONTS_DIR, "LICENSE.txt")),
          "assets/fonts/LICENSE.txt")
    check("font-dele indlejret i pakken", len(parts) == len(src) and parts,
          "%d dele" % len(parts))
    matched = [src.get(hashlib.sha256(z.read(n)).hexdigest(), "?" + n) for n in parts]
    check("hver del er byte-identisk med sin kildefil", not any(m.startswith("?") for m in matched),
          ", ".join(sorted(matched)))

    rels = etree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
    FONT_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
    font_rels = {r.get("Id"): r.get("Target")
                 for r in rels.findall("pr:Relationship", NS) if r.get("Type") == FONT_RT}
    check("font-relationships", len(font_rels) == len(parts), " ".join(sorted(font_rels)))

    pres = etree.fromstring(z.read("ppt/presentation.xml"))
    kids = [etree.QName(k).localname for k in pres if isinstance(k.tag, str)]
    check("p:embeddedFontLst findes", "embeddedFontLst" in kids)
    idx = [ORDER.index(k) for k in kids if k in ORDER]
    check("elementrækkefølge følger ECMA-376", idx == sorted(idx),
          " -> ".join(k for k in kids if k in ORDER))

    lst = pres.find("p:embeddedFontLst", NS)
    for ef in (lst.findall("p:embeddedFont", NS) if lst is not None else []):
        font = ef.find("p:font", NS)
        tf = font.get("typeface")
        sub = [etree.QName(c).localname for c in ef]
        check("%s: barn-rækkefølge font,regular,bold" % tf,
              sub[:1] == ["font"] and sub[1:] == sorted(sub[1:], key=
                  lambda s: ["regular", "bold", "italic", "boldItalic"].index(s)),
              ",".join(sub))
        check("%s: panose + pitchFamily" % tf,
              bool(font.get("panose")) and bool(font.get("pitchFamily")),
              "panose=%s pitch=%s" % (font.get("panose"), font.get("pitchFamily")))
        for slot in sub[1:]:
            rid = ef.find("p:" + slot, NS).get("{%s}id" % NS["r"])
            tgt = font_rels.get(rid, "")
            check("%s: %s -> %s" % (tf, slot, tgt),
                  rid in font_rels and ("ppt/" + tgt.lstrip("/")) in names, rid)


# ============================================================== 3. HTML/PPTX-PARITET
def _strip_comments(css):
    """Fjern /* ... */ — ellers bliver kommentartekst læst som selector."""
    return re.sub(r"/\*.*?\*/", "", css, flags=re.S)


def _css_block(css, selector):
    """Returnér erklæringerne i den FØRSTE regel med præcis denne selector."""
    css = _strip_comments(css)
    for m in re.finditer(r"([^{}]+)\{([^{}]*)\}", css):
        # alt før sidste ';' er @import/@charset-sætninger, ikke selector
        raw = m.group(1).split(";")[-1]
        sels = [" ".join(s.split()) for s in raw.split(",")]
        if selector in sels:
            return m.group(2)
    return None


def _css_px(css, selector, prop):
    block = _css_block(css, selector)
    if block is None:
        return None
    m = re.search(r"(?m)(?:^|;)\s*%s\s*:\s*([^;]+)" % re.escape(prop), block)
    if not m:
        return None
    nums = re.findall(r"(-?\d+(?:\.\d+)?)px", m.group(1))
    return float(nums[0]) if nums else None


def _css_num(css, selector, prop):
    """Som _css_px, men til enhedsløse værdier (fx line-height: 1.05)."""
    block = _css_block(css, selector)
    if block is None:
        return None
    m = re.search(r"(?m)(?:^|;)\s*%s\s*:\s*([^;]+)" % re.escape(prop), block)
    if not m:
        return None
    n = re.findall(r"-?\d+(?:\.\d+)?", m.group(1))
    return float(n[0]) if n else None


def _font_metrics(path):
    """(unitsPerEm, ascender, descender) fra en TTF's head/hhea-tabeller."""
    import struct
    try:
        b = open(path, "rb").read()
        num = struct.unpack(">H", b[4:6])[0]
        tabs = {}
        for i in range(num):
            o = 12 + i * 16
            tabs[b[o:o + 4]] = struct.unpack(">I", b[o + 8:o + 12])[0]
        upem = struct.unpack(">H", b[tabs[b"head"] + 18:tabs[b"head"] + 20])[0]
        asc, desc = struct.unpack(">hh", b[tabs[b"hhea"] + 4:tabs[b"hhea"] + 8])
        return upem, asc, desc
    except Exception:
        return None


def check_parity():
    import build_pptx as bp
    css = open(os.path.join(HERE, "slides.css"), encoding="utf-8").read()

    def inch_px(emu):
        return emu / 914400.0 * PX_PER_INCH

    pairs = [
        # (label, css-værdi, pptx-px, tolerance)
        ("sikker margin", _css_px(css, ":root", "--slide-pad"), inch_px(bp.MARGIN), 1),
        ("logo: luft til højre", _css_px(css, ".slide::after", "right"),
         inch_px(bp.LOGO_INSET), 1),
        ("logo: luft til bunden", _css_px(css, ".slide::after", "bottom"),
         inch_px(bp.LOGO_INSET), 1),
        ("logo: bredde", _css_px(css, ".slide::after", "width"), inch_px(bp.LOGO_W), 2),
        ("forside: kassens bredde", _css_px(css, ".slide--title .frame", "width"),
         inch_px(bp.COVER_W), 1),
        ("forside: højde uden manchet",
         _css_px(css, ".slide--title .frame", "min-height"), inch_px(bp.COVER_H_SHORT), 1),
        ("forside: højde med manchet",
         _css_px(css, ".slide--title .frame:has(.lead)", "min-height"),
         inch_px(bp.COVER_H), 1),
    ]
    for label, got, want, tol in pairs:
        if got is None:
            check("paritet: %s" % label, False, "kunne ikke læses i slides.css")
            continue
        check("paritet: %s" % label, near(got, want, tol),
              "css %gpx  vs  pptx %.1fpx" % (got, want))

    # --- accent-bjælkens afstand under overskriftens baseline ---
    # Den vigtigste paritet, og den sværeste at se: de to formater sætter
    # overskriften forskelligt (PPTX-tekstboks vs. CSS-linjeboks), så man kan
    # ikke bare sammenligne absolutte positioner. Vi regner baselinen ud i
    # begge formater ud fra fontens egne metrikker og sammenligner afstanden.
    m = _font_metrics(os.path.join(FONTS_DIR, "IBMPlexSans-Bold.ttf"))
    if m is None:
        check("paritet: accent under baselinen", False, "kunne ikke læse fontmetrikker")
        return
    upem, asc, desc = m

    # PPTX: tekstboksens top + python-pptx' default tIns (0.05") + ascender
    T_INS = 0.05
    pptx_baseline = bp.TITLE_TOP / 914400.0 + T_INS + asc / upem * bp.TITLE_PT / 72.0
    pptx_gap = (bp.ACCENT_Y / 914400.0 - pptx_baseline) * PX_PER_INCH

    # HTML: linjeboksens halve leading + ascender. Slide-padding går ud med
    # sig selv, fordi begge sider måles fra overskriftens egen top.
    fs = _css_px(css, ".slide-title", "font-size")
    lh = _css_num(css, ".slide-title", "line-height")
    mt = _css_px(css, ".slide-title::after", "margin")
    if None in (fs, lh, mt):
        check("paritet: accent under baselinen", False,
              "fandt ikke font-size/line-height/margin i slides.css")
        return
    line_box = lh * fs
    content_area = (asc - desc) / upem * fs
    half_leading = (line_box - content_area) / 2.0
    html_gap = line_box + mt - half_leading - asc / upem * fs

    check("paritet: accent under baselinen", near(html_gap, pptx_gap, 1.5),
          "css %.1fpx  vs  pptx %.1fpx" % (html_gap, pptx_gap))

    # kassens top: slide-padding + 3px ramme + frame margin-top
    mt = _css_px(css, ".slide--title .frame", "margin-top")
    pad = _css_px(css, ":root", "--slide-pad")
    if mt is not None and pad is not None:
        got = mt + pad + 3
        check("paritet: forside: kassens top", near(got, inch_px(bp.COVER_TOP), 2),
              "css %gpx  vs  pptx %.1fpx" % (got, inch_px(bp.COVER_TOP)))


# ============================================================== 4. GALLERI-DÆKNING
def check_gallery_coverage():
    src = open(os.path.join(HERE, "build_pptx.py"), encoding="utf-8").read()
    example = open(os.path.join(HERE, "deck.example.py"), encoding="utf-8").read()
    methods = set(re.findall(r"(?m)^    def ([a-z][a-z_0-9]*)\(self", src)) - {"save"}
    used = set(re.findall(r"deck\.([a-z_0-9]+)\(", example))
    missing = sorted(methods - used)
    check("galleriet viser hver slide-type", not missing,
          "%d typer%s" % (len(methods), "" if not missing else " — mangler: " + ", ".join(missing)))


# ==================================================================== 5. OPRYDNING
def check_cleanup():
    tpl = open(os.path.join(HERE, "template.html"), encoding="utf-8").read()
    # kicker-eyebrows må ikke stå umiddelbart over en overskrift
    over_title = re.findall(
        r'<p class="kicker">[^<]*</p>\s*<h[12] class="slide-(?:display|title)"', tpl)
    check("ingen kicker-label over en overskrift i galleriet", not over_title,
          "%d fundet" % len(over_title))
    check("galleriet linker slides.css (ikke det udfasede shim)",
          'href="slides.css"' in tpl and "slides-oes.css" not in tpl)
    example = open(os.path.join(HERE, "deck.example.py"), encoding="utf-8").read()
    check("eksemplet bruger ikke kicker= eller theme=",
          "kicker=" not in example and "theme=" not in example)


# ========================================================================== MAIN
def main():
    print(__doc__.strip().split("\n\n")[0])
    path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "example.pptx")

    if len(sys.argv) == 1:
        head("Bygger galleriet")
        r = subprocess.run([sys.executable, os.path.join(HERE, "deck.example.py")],
                           cwd=HERE, capture_output=True, text=True)
        sys.stdout.write(r.stdout)
        sys.stderr.write(r.stderr)
        if r.returncode != 0:
            print("\nFEJL: deck.example.py fejlede — kan ikke fortsætte.")
            return 1

    if not os.path.exists(path):
        print("\nFEJL: %s findes ikke. Kør 'python deck.example.py' først." % path)
        return 1

    sys.path.insert(0, HERE)

    head("0. Referencemål (udmålt i de to referencedeck)")
    check_reference_measurements()

    head("1. Geometri i %s" % os.path.basename(path))
    n = check_pptx_geometry(path)
    print("     (%d slides)" % n)

    head("2. Font-indlejring")
    check_font_embedding(path)

    head("3. HTML/PPTX-paritet")
    check_parity()

    head("4. Galleri-dækning")
    check_gallery_coverage()

    head("5. Oprydning")
    check_cleanup()

    print()
    if _failures:
        print("%d TJEK FEJLEDE:" % len(_failures))
        for f in _failures:
            print("  - " + f)
        return 1
    print("ALT BESTAAET")
    return 0


if __name__ == "__main__":
    sys.exit(main())
